"""Create an editable project introduction and a layout preview from local sources.

Build dependencies: python-pptx and Pillow, installed in .tools/presentation-libs.
"""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.tools/presentation-libs'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

OUT = ROOT / 'docs/presentations'
OUT.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
prs.core_properties.title = 'DeskPet｜把信息变成陪伴'
prs.core_properties.subject = '评比展示 · 主动热点搜索与播报 · 个性化陪伴 · 技术实现'
prs.core_properties.author = 'DeskPet'
prs.core_properties.keywords = 'DeskPet,桌宠,Rust,Tauri,Android,Edge TTS'
BG, PANEL, INK, MUTED, TEAL, CORAL, WHITE = '#F6F3EC', '#FFFFFF', '#163A3B', '#607474', '#1F7771', '#E87E58', '#FFFFFF'
FONT = 'Microsoft YaHei'
FONTFILE = Path('C:/Windows/Fonts/msyh.ttc')
previews = []
slide = canvas = draw = None

def rgb(color):
    return RGBColor.from_string(color.lstrip('#'))

def rect(x, y, w, h, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(color); s.line.fill.background()
    if rounded:
        s.adjustments[0] = .12
        draw.rounded_rectangle((x, y, x+w, y+h), radius=18, fill=color)
    else:
        draw.rectangle((x, y, x+w, y+h), fill=color)
    return s

def text(x, y, w, h, value, size=28, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    font = ImageFont.truetype(str(FONTFILE), round(size * 100/72))
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.line_spacing = 1.15
        draw.text((x, y+i*size*100/72*1.15), line, font=font, fill=color)
    return box

def picture(path, x, y, w, h, crop=None):
    s = slide.shapes.add_picture(str(path), Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100))
    im = Image.open(path).convert('RGBA')
    if crop:
        l,t,r,b = crop
        s.crop_left, s.crop_top, s.crop_right, s.crop_bottom = crop
        im = im.crop((int(im.width*l), int(im.height*t), int(im.width*(1-r)), int(im.height*(1-b))))
    im = im.resize((round(w), round(h)), Image.Resampling.NEAREST if 'sprites' in str(path) else Image.Resampling.LANCZOS)
    canvas.paste(im, (round(x), round(y)), im)

def pet(state, x, y, size=220):
    picture(ROOT / f'crates/deskpet-app/frontend/sprites/{state}.png', x,y,size,size, (0,0,.875,0))

def new(title, subtitle='', dark=False, sources='README.md; docs/product-features.md'):
    global slide, canvas, draw
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = INK if dark else BG
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(background)
    canvas = Image.new('RGB', (1600,900), background); draw = ImageDraw.Draw(canvas)
    previews.append(canvas)
    rect(60,56,44,6,CORAL)
    text(122,37,1000,30,'DESKPET  /  PROJECT INTRODUCTION',13, '#B8CFCA' if dark else MUTED,True)
    if title:
        text(60,112,1480,90,title,34,WHITE if dark else INK,True)
    if subtitle:
        text(60,203,1480,60,subtitle,17,'#B8CFCA' if dark else MUTED)
    text(60,852,1250,25,'DeskPet  ·  项目介绍  ·  2026.10.24',11,'#B8CFCA' if dark else MUTED)
    text(1460,847,80,30,f'{len(prs.slides):02d}',14,CORAL,True)
    slide.notes_slide.notes_text_frame.text = '资料依据（项目相对路径）：'+sources+'\n内容以 2026-10-24 项目源码与本次会话已交付功能为基线。示意图不代表实机截图。'

def card(x,y,w,h,number,title,body):
    rect(x,y,w,h,PANEL,True)
    text(x+25,y+20,w-50,40,number,20,CORAL,True)
    text(x+25,y+78,w-50,55,title,25,INK,True)
    text(x+25,y+150,w-50,h-155,body,18,MUTED)

new('',dark=True)
text(65,200,850,85,'DeskPet',60,WHITE,True)
text(70,320,1000,90,'把信息变成陪伴',42,WHITE,True)
text(74,450,1000,125,'主动搜索热点，主动分享你关心的新鲜事\n兴趣驱动 × 个性表达 × 语音陪伴',23,'#B8CFCA')
rect(75,650,420,64,TEAL,True); text(100,662,380,45,'主动分享  /  个性表达  /  轻量互动',16,WHITE,True)
rect(1130,230,355,430,'#214B4A',True); pet('happy',1155,295,300)
text(1185,647,300,40,'让桌面多一点回应',18,'#B8CFCA')

new('01  同一套核心，两种陪伴形态','Windows 桌面宠物与 Android 横屏界面，共享热点采集、画像和内容生成能力。',sources='docs/product-features.md; docs/design/desktop-companion.md')
rect(60,295,710,455,WHITE,True); rect(805,295,735,455,WHITE,True)
text(90,320,650,50,'Windows｜桌面上的伙伴',26,INK,True)
pet('idle',95,400,215)
rect(345,408,365,120,BG,True); text(366,425,325,100,'主动分享新鲜事\n气泡展示 + 语音朗读',18)
text(95,637,620,90,'透明窗口 · 可拖动 · 托盘设置\n阅读时暂停关闭 · 反馈学习',19,MUTED)
text(840,320,660,50,'Android｜头部屏幕界面',26,INK,True)
rect(850,405,640,190,INK,True); pet('curious',870,421,165)
rect(1100,427,355,105,'#FFF0CC',True); text(1120,442,315,80,'右上角消息气泡\n正文可独立滚动',17)
text(850,637,640,95,'横屏沉浸 · 前台亮屏 · 触摸互动\n待机宠物居中 · 消息气泡位于右上方',18,MUTED)

new('02  核心闭环：主动发现，主动播报','不等待用户提问：主动采集 → 兴趣筛选 → 择机分享 → 反馈学习。',sources='crates/deskpet-app/src/runtime.rs; crates/deskpet-scheduler/src/scheduler.rs; docs/architecture.md')
steps=[('主动搜索热点','主动关注多平台\n发现新鲜内容'),('判断是否分享','综合兴趣与热度\n结合间隔和勿扰'),('生成宠物表达','排除已播报内容\n融入人设与画像'),('主动展示播报','气泡呈现新鲜事\nWindows 语音朗读'),('反馈更新偏好','喜欢 / 不感兴趣\n用于后续内容筛选')]
for i,(title,body) in enumerate(steps):
    x=60+i*301
    card(x,325,275,310,f'0{i+1}',title,body)
    if i<4: text(x+280,447,25,40,'›',26,CORAL,True)
rect(60,700,1480,94,INK,True)
text(87,719,1425,65,'宠物持续关注新内容，在合适的时机主动分享，让陪伴自然发生。',22,WHITE)

new('03  内容如何变得更贴近用户','以可解释的规则组合兴趣、热度、新鲜度与反馈。',sources='docs/architecture.md; crates/deskpet-agent/src/scoring.rs')
rect(60,300,930,445,WHITE,True)
text(90,326,860,55,'内容处理链',26,INK,True)
text(90,413,850,255,'目标支持平台：知乎 / B站 / 微博 / 抖音 / 小红书\n→ 标签处理、可选向量与去重入库\n→ 兴趣匹配、热度与新鲜度综合评分\n→ 宠物人设 + 模型表达 / 无 Key 模板降级',21)
rect(1040,300,500,445,INK,True)
text(1070,333,440,60,'反馈让偏好持续变化',25,WHITE,True)
text(1070,433,440,205,'喜欢 → 调整标签兴趣\n不感兴趣 → 降低相关偏好\n定期维护 → 兴趣权重衰减',21,'#B8CFCA')
text(70,785,1450,40,'信息经过采集、筛选、记录与反馈，形成持续运转的个性化内容链路。',18,TEAL,True)

new('04  个性化表达：同样的热点，不同的陪伴','宠物名称、性格、表达偏好与背景参与模型回复，六种动画情绪增强互动感。')
states=[('idle','待机'),('happy','开心'),('curious','好奇'),('excited','兴奋'),('sleepy','困倦'),('bored','无聊')]
for i,(state,label) in enumerate(states):
    x=60+i*250
    rect(x,320,230,265,WHITE,True); pet(state,x+30,345,170); text(x+70,527,150,40,label,21,INK,True)
rect(60,650,1480,125,INK,True)
text(90,675,1420,90,'内容有来源，表达有性格，互动有情绪。\n用户既可以定义宠物如何说话，也可以通过反馈告诉它自己喜欢什么。',22,WHITE)

new('05  Rust 核心复用，平台界面分层','Rust workspace + Tauri + HTML/CSS/JavaScript + SQLite。',sources='docs/project-structure.md; docs/architecture.md')
for x,w,title,body in [(60,710,'Windows 平台层','透明桌宠 / 独立气泡 / 托盘 / 语音进程'),(805,735,'Android 平台层','Activity / 横屏沉浸 / 页面内气泡 / 前台亮屏')]:
    rect(x,290,w,125,INK,True); text(x+25,305,w-50,45,title,24,WHITE,True); text(x+25,363,w-50,40,body,17,'#B8CFCA')
rect(60,450,1480,105,TEAL,True); text(90,477,1410,65,'deskpet-app  ·  配置、IPC、生命周期与共享前端',26,WHITE,True)
for i,(title,body) in enumerate([('collector','多平台采集'),('agent','筛选与表达'),('scheduler','主动调度与策略'),('tts','合成与播放')]):
    x=60+i*378; rect(x,590,346,105,WHITE,True); text(x+20,603,306,40,title,23,TEAL,True); text(x+20,651,306,40,body,17,MUTED)
rect(60,730,1480,70,WHITE,True); text(85,743,1420,50,'core 共享类型    +    storage / SQLite    +    本地配置与轮转日志',23,INK,True)

new('06  Windows 语音：先看到，再听到','Edge TTS 在线合成，便携 Python 执行 worker，mpv 播放本地音频。',sources='docs/windows-edge-integration-plan.md; crates/deskpet-app/src/windows_speech.rs; crates/deskpet-tts/src/lib.rs')
for i,(title,body) in enumerate([('气泡确认显示','正文先渲染\n合成不阻塞显示'),('异步合成','Edge TTS 生成 MP3\n失败保留文字'),('播放并保持气泡','合成与播放期间\n暂停自动关闭'),('结束或主动停止','结束恢复阅读计时\n关闭气泡即停止')]):
    card(60+i*378,310,346,315,f'0{i+1}',title,body)
rect(60,690,1480,115,INK,True)
text(90,711,1420,90,'音色自由选择，语音开关由用户掌控；关闭气泡即可停止对应朗读。\n链接自动结束朗读，新消息替换旧任务，让内容呈现与声音保持同步。',22,WHITE)

new('07  Android：为桌面陪伴设计的横屏体验','界面布局示意：待机时宠物居中，消息出现后让出阅读区域。',sources='docs/design/desktop-companion.md; crates/deskpet-app/frontend/mobile.css')
rect(60,300,950,450,INK,True)
text(95,324,840,40,'DeskPet                              横屏 · 沉浸 · 前台亮屏',17,'#B8CFCA')
pet('happy',100,423,265)
rect(420,390,535,230,'#FFF0CC',True)
text(450,415,470,175,'发现一条你可能感兴趣的热点\n\n消息在右上角呈现\n完整正文在气泡内滚动',23,INK)
text(1080,325,440,60,'面向小屏陪伴场景',25,INK,True)
text(1080,425,450,270,'横屏布局便于桌面观看\n沉浸界面突出宠物与消息\n前台亮屏保持陪伴感\n气泡内部滚动阅读长文\n触摸反馈与双击辅助交互\n动画表达当前情绪',21,MUTED)

new('08  实体桌宠概念：让软件拥有一个陪伴载体','设计概念展示：以 Android 横向屏幕作为头部，将信息互动融入桌面空间。',sources='docs/design/deskpet-physical-concept.png; docs/design/desktop-companion.md')
picture(ROOT/'docs/design/deskpet-physical-concept.png',60,277,790,527)
text(920,310,620,65,'把陪伴放在桌面上',30,INK,True)
text(920,425,620,260,'横向屏幕承载宠物表情\n圆润机身营造亲近感\n消息气泡连接信息与情绪\n自然融入学习、办公与生活桌面',22,MUTED)
text(920,736,620,75,'形态设计概念图\n软件能力与外观设计相互呼应',16,TEAL)

new('09  六类设置，让陪伴适应不同用户','Windows 设置提供完整配置入口，将内容、节奏、表达与声音交给用户。',sources='crates/deskpet-app/frontend/settings.html; crates/deskpet-app/frontend/settings.js')
settings=[('常规策略','勿扰时段 / 播报间隔\n播报评分阈值'),('宠物画像','名称 / 性格 / 表达偏好\n身份与背景设定'),('模型服务','API 地址 / Key / 模型名称\n可选向量模型'),('热点采集','平台启用 / 独立采集周期\n默认采集周期 / 来源配置'),('语音播报','语音开关 / 中文音色\n晓晓 / 云希 / 晓伊 / 云健'),('用户画像','人物描述 / 兴趣主题\n兴趣权重可编辑')]
for i,(title,body) in enumerate(settings):
    x=60+(i%3)*505; y=285+(i//3)*245
    rect(x,y,470,220,WHITE,True)
    text(x+24,y+20,425,45,title,25,INK,True)
    text(x+24,y+90,425,110,body,20,MUTED)
text(75,805,1450,40,'可视化配置 + 本地保存：同一套应用可以呈现不同的陪伴风格。',18,TEAL,True)

new('10  配置示例：同一只宠物，不同的陪伴节奏','以下为可配置项的场景示例，用户可以自行组合，并非固定预设。',sources='crates/deskpet-app/frontend/settings.html; crates/deskpet-app/src/config.rs')
card(60,300,470,425,'学习场景','安静而有选择地分享','兴趣：人工智能 / 科技\n间隔：大 · 60 分钟\n阈值：更高，更挑选内容\n人设：简洁、温和、重点明确\n语音：按需开启')
card(565,300,470,425,'日常场景','让新鲜事自然到来','兴趣：科技 / 游戏 / 生活\n间隔：中 · 45 分钟\n阈值：适度分享\n人设：活泼、轻松、有好奇心\n语音：晓晓 / 云希等可选')
card(1070,300,470,425,'专注时段','用户决定何时被打扰','设置自己的勿扰时段\n按平台调整采集周期\n关闭或开启语音播报\n调整兴趣主题与权重\n通过反馈持续更新偏好')
text(80,770,1440,45,'不是统一推送给所有人，而是围绕每个人的兴趣、节奏和表达偏好进行配置。',20,TEAL,True)

new('11  项目亮点：把信息服务做成主动陪伴','从内容获取到桌面表达，形成一套可运行、可配置、可持续反馈的产品闭环。')
card(60,300,470,390,'主动性','主动发现与分享','主动搜索多平台热点\n结合策略主动选择分享时机\n气泡与 Windows 语音主动呈现')
card(565,300,470,390,'个性化','兴趣驱动的表达','兴趣画像参与内容筛选\n宠物人设影响表达方式\n反馈改变后续偏好权重')
card(1070,300,470,390,'产品化','完整应用与交付','桌宠、托盘、气泡与设置\nRust 共享核心与双平台界面\n本地数据持久化与 Release 交付')
text(80,750,1450,60,'主动发现 + 个性筛选 + 宠物表达 + 语音陪伴',29,TEAL,True)

new('',dark=True)
text(75,180,1380,90,'让信息有用，让陪伴有回应。',43,WHITE,True)
text(80,335,1250,155,'DeskPet 主动搜索热点，筛选你关心的内容，\n再以宠物的表情、口吻和声音，主动与你分享。',27,'#B8CFCA')
rect(80,575,1000,105,TEAL,True); text(110,601,940,70,'让宠物主动关注世界，也主动关心你的兴趣。',25,WHITE,True)
pet('happy',1230,530,250)

path = OUT/'DeskPet-项目介绍-评比版.pptx'
prs.save(path)
# Preview is generated from the same layout primitives; PowerPoint is the final renderer.
thumbs=[]
for i,im in enumerate(previews,1):
    im.save(OUT/f'slide-{i:02d}.png')
    thumb=im.resize((480,270),Image.Resampling.LANCZOS)
    thumbs.append(thumb)
sheet=Image.new('RGB',(1480, len(range(0,len(thumbs),3))*300+20),'#DDE4DF')
for i,im in enumerate(thumbs): sheet.paste(im,(10+(i%3)*490,10+(i//3)*300))
sheet.save(OUT/'DeskPet-项目介绍-评比版-总览.png')
readback=Presentation(path)
assert len(readback.slides)==13
for i,s in enumerate(readback.slides,1):
    for shape in s.shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=prs.slide_width+Inches(.02), (i,shape.name)
        assert shape.top+shape.height<=prs.slide_height+Inches(.02), (i,shape.name)
print(f'Created and validated {len(readback.slides)} slides: {path}')
print(f'Layout preview: {OUT / "DeskPet-项目介绍-评比版-总览.png"}')
